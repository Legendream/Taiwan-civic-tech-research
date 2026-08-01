# -*- coding: utf-8 -*-
"""
08_update_old_report.py — 把舊報告《選題依據與建議報告》的數字更新到 N=127。

做法：**原地替換數字，完全不動 Claire 寫的文字與版面**，並逐筆記錄
「改哪裡、從什麼改成什麼」，讓她能逐條核對而不必重讀整份報告。

分母慣例（很重要）：
  舊報告用「總回收份數」與「該層總人數」當分母
  （查核：群眾協力 74/121 = 61.2%、政府預算中間層 31/50 = 62%，與舊報告完全吻合）。
  v2 **沿用同一慣例**，這樣 v2 與 v1 才能直接對照。
  （新的《生態系分析報告》用較嚴謹的「該題有作答者」當分母，兩者差異 ≤1 人、≤0.8pt，
   已在變更紀錄中標明。）

輸出：公民科技指引_選題依據與建議報告_v2.docx
      分析結果/08_舊報告變更紀錄.csv

執行：python3 08_update_old_report.py
"""

import shutil

import docx
import pandas as pd

import common as C

SRC = C.OUT_DIR / "公民科技指引_選題依據與建議報告.docx"
DST = C.OUT_DIR / "公民科技指引_選題依據與建議報告_v2.docx"

CHANGES = []


def log(where, before, after, note=""):
    if str(before).strip() != str(after).strip():
        CHANGES.append({"位置": where, "原本": before, "改為": after, "說明": note})


def set_cell(cell, text, where, note=""):
    """
    替換儲存格文字，保留第一個 run 的格式。

    表格儲存格放的是單一數值或短標籤、格式一致，併進 runs[0] 不會像段落那樣
    造成格式混亂；但仍檢查一次，若儲存格內有格式不一致的多個 run 就示警。
    """
    before = cell.text.strip()
    if before == str(text).strip():
        return
    p = cell.paragraphs[0]
    if len({(r.bold, r.italic, r.underline) for r in p.runs}) > 1:
        print(f"    ⚠ {where}：儲存格含格式不一致的多個 run，併入首個 run 會統一格式")
    log(where, before, text, note)
    if p.runs:
        p.runs[0].text = str(text)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(str(text))
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def replace_in_paragraph(p, pairs, where):
    """
    在段落中做字串替換，**逐 run 就地替換**以保留行內格式。

    ⚠ 不可把整段文字併進 runs[0] 再清空其他 run：
      Word 的粗體／字級是掛在 run 上的，一併進 runs[0] 會讓整段繼承 runs[0] 的格式。
      先前就是這樣把 P13、P50 整段變成粗體（那兩段的 run0 剛好是粗體標題片語）。
    """
    before_text = p.text
    changed_any = False
    for a, b in pairs:
        for r in p.runs:
            if a in r.text:
                r.text = r.text.replace(a, b)
                changed_any = True
    if changed_any:
        log(where, before_text, p.text)
        return
    # 若目標字串跨越 run 邊界，逐 run 替換會抓不到；此時明確報出來，不要靜默略過
    fallback = before_text
    for a, b in pairs:
        fallback = fallback.replace(a, b)
    if fallback != before_text:
        raise RuntimeError(
            f"{where}：待替換字串跨越 run 邊界，逐 run 替換抓不到。\n"
            f"  原文：{before_text[:80]}\n"
            f"  請改用更短、不跨 run 的比對字串，或人工處理這一段。")


def replace_embedded_figure(docx_path, old_png, new_png):
    """
    在 docx（本質是 zip）裡就地替換一張內嵌圖片的位元組。

    以「舊圖的位元組」比對來定位，不靠檔名雜湊硬編碼——這樣圖片在文件中換位置、
    或 Word 重新命名 media 檔時仍然找得到。新舊圖長寬比必須相同，否則會被拉伸。
    """
    import shutil
    import struct
    import tempfile
    import zipfile

    def png_size(path):
        with open(path, "rb") as f:
            f.read(16)
            return struct.unpack(">II", f.read(8))

    old_bytes = old_png.read_bytes()
    new_bytes = new_png.read_bytes()
    ow, oh = png_size(old_png)
    nw, nh = png_size(new_png)
    assert abs(ow / oh - nw / nh) < 0.01, \
        f"新舊圖長寬比不同（{ow}x{oh} vs {nw}x{nh}），換進去會被拉伸"

    with zipfile.ZipFile(docx_path) as z:
        names = z.namelist()
        target = [n for n in names if n.startswith("word/media/") and z.read(n) == old_bytes]
        if not target:
            print("    ⚠ 找不到與舊圖位元組相符的內嵌圖片，未替換（圖可能已經換過或被壓縮）")
            return
        items = [(n, z.read(n)) for n in names]

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    tmp.close()
    with zipfile.ZipFile(tmp.name, "w", zipfile.ZIP_DEFLATED) as out:
        for n, data in items:
            out.writestr(n, new_bytes if n in target else data)
    shutil.move(tmp.name, docx_path)
    log(f"內嵌圖 {target[0]}", "狩野四象限說明圖（N=120 版）",
        "狩野四象限說明圖（N=126 版）", "圖 3-1，長寬比不變")
    print(f"    ✓ 已替換內嵌圖 3-1（{target[0]}），新圖 {nw}x{nh} 與原圖同長寬比")


def main():
    if not SRC.exists():
        raise SystemExit(f"找不到舊報告：{SRC}")
    shutil.copy(SRC, DST)
    d = docx.Document(DST)

    new = C.load_clean()
    N = len(new)                                   # 127
    layer_n = new[C.LAYER].value_counts().to_dict()

    kano_all = pd.read_csv(C.TABLE_DIR / "02_狩野_全體.csv").set_index("主題")
    kano_lyr = pd.read_csv(C.TABLE_DIR / "02_狩野_三層.csv")
    rank = pd.read_csv(C.TABLE_DIR / "07_舊報告差異_資源短文排序.csv")

    # ---------- 段落數字 ----------
    para_pairs = [
        ("N≈120", "N≈126"), ("N=120", "N=126"), ("回收 121 份", "回收 127 份"),
        ("回收總數：121 份", "回收總數：127 份"), ("有效填答 120", "有效填答 126"),
        ("填答 121", "填答 127"), ("填答 120", "填答 126"),
        ("N=24", "N=26"), ("（回收 121 份）", "（回收 127 份）"),
        # 內文以中文敘述提到的層人數，也要一起更新，否則 v2 會自相矛盾。
        # ⚠ 一律用帶前後文的錨定字串。先前用裸的 ("24 人","26 人")、("N=50","N=54")，
        #    會把文件裡任何一個「24 人」都改掉——例如某活動出席 24 人這種無關的數字。
        ("那 24 人", "那 26 人"),
        ("層，N=50）", "層，N=54）"),      # P44：「即『接觸未參與』層，N=50）為主要目標讀者」
    ]

    # 替換前先確認每個錨定字串在文件中真的存在，否則就是錨定寫錯、會靜默漏改
    doc_text = "\n".join(p.text for p in d.paragraphs)
    for a, _ in para_pairs:
        if a not in doc_text:
            print(f"    ⚠ 錨定字串未在文件中出現，將不會有任何替換：{a!r}")
    for i, p in enumerate(d.paragraphs):
        replace_in_paragraph(p, para_pairs, f"段落 P{i}")

    # ---------- 表1：入坑三層 N ----------
    t = d.tables[1]
    for row, layer in zip(t.rows[1:], [C.L_NEVER, C.L_AWARE, C.L_DONE]):
        set_cell(row.cells[2], layer_n[layer], f"表1 {layer} N")

    # ---------- 表3：資源短文整合排序（含名次移動，整表重寫）----------
    t = d.tables[3]
    top = rank.sort_values("排名_新").head(8)
    for i, (_, r) in enumerate(top.iterrows(), start=1):
        row = t.rows[i]
        set_cell(row.cells[0], int(r["排名_新"]), f"表3 第{i}列 排名")
        set_cell(row.cells[1], r["類型"], f"表3 第{i}列 類型")
        set_cell(row.cells[2], SHORT.get(r["選項"], r["選項"]), f"表3 第{i}列 主題")
        set_cell(row.cells[3], f"{r['比例_新']:.1%}", f"表3 第{i}列 勾選率")
    ninth = rank[rank["排名_新"] == 9].iloc[0]
    gap = top["比例_新"].min() - ninth["比例_新"]
    set_cell(t.rows[9].cells[2],
             f"— 前 8 名分界線（第 9 名 {ninth['比例_新']:.1%}，有 {gap * 100:.0f} 個百分點斷層）—",
             "表3 分界線說明")
    for i, (_, r) in enumerate(rank[rank["排名_新"].isin([9, 10])].iterrows(), start=10):
        row = t.rows[i]
        set_cell(row.cells[0], int(r["排名_新"]), f"表3 第{i}列 排名")
        set_cell(row.cells[1], r["類型"], f"表3 第{i}列 類型")
        set_cell(row.cells[2], SHORT.get(r["選項"], r["選項"]), f"表3 第{i}列 主題")
        set_cell(row.cells[3], f"{r['比例_新']:.1%}", f"表3 第{i}列 勾選率")

    # ---------- 表5：中間層狩野 SI ----------
    mid = kano_lyr[kano_lyr["分群"] == C.L_AWARE].set_index("主題")
    for row in d.tables[5].rows[1:]:
        topic = row.cells[1].text.strip()
        full = MATCH.get(topic)
        if full and full in mid.index:
            r = mid.loc[full]
            set_cell(row.cells[2], f"{r['SI']:.2f} / {r['落點'].split()[-1]}",
                     f"表5 {topic} 中間層 SI", "分母＝接觸未參與層有效樣本")

    # ---------- 表6：中間層勾選率 ----------
    mid_n = layer_n[C.L_AWARE]
    picks = {}
    for fname in ["03_三層×議題領域.csv", "03_三層×跨領域工具.csv"]:
        t6 = pd.read_csv(C.TABLE_DIR / fname)
        for _, r in t6[t6["分群"] == C.L_AWARE].iterrows():
            picks[r["選項"]] = r["分子"] / mid_n          # 沿用舊慣例：分母＝該層總人數
    # 表6 是「依中間層勾選率排序」的建議表：數字更新後，列的順序也必須跟著重排，
    # 否則會出現 41% 排在三個 39% 下面這種自相矛盾。
    # 做法：保留每一列原本的「建議」標記（★／△ 是 Claire 的編輯判斷，不動），
    #       只在同一個建議層級內，依新的勾選率重新排序。
    t6 = d.tables[6]
    body = []
    for row in t6.rows[1:]:
        topic = row.cells[2].text.strip()
        full = MATCH.get(topic.replace("（三選二）", ""))
        body.append({
            "建議": row.cells[0].text.strip(), "類型": row.cells[1].text.strip(),
            "主題": topic, "率": picks.get(full, -1),
        })
    order = {"★": 0, "△": 1, "○": 2, "▽": 3}
    body.sort(key=lambda r: (order.get(r["建議"], 9), -r["率"]))
    for row, b in zip(t6.rows[1:], body):
        set_cell(row.cells[0], b["建議"], "表6 建議標記（重排後）")
        set_cell(row.cells[1], b["類型"], "表6 類型（重排後）")
        set_cell(row.cells[2], b["主題"], "表6 主題（重排後）")
        set_cell(row.cells[3], f"{b['率']:.0%}",
                 f"表6 {b['主題']} 中間層勾選率", f"分母＝{mid_n}（接觸未參與層人數）")

    # ---------- 表7：全體狩野摘要 ----------
    for row in d.tables[7].rows[1:]:
        full = MATCH.get(row.cells[0].text.strip())
        if full and full in kano_all.index:
            r = kano_all.loc[full]
            for ci, val in [(1, int(r["A"])), (2, int(r["O"])), (3, int(r["M"])),
                            (4, int(r["I"])), (5, f"{r['SI']:.2f}"), (6, f"{r['DSI']:.2f}")]:
                set_cell(row.cells[ci], val, f"表7 {row.cells[0].text.strip()} 第{ci}欄")

    # ---------- 表8／表9：三層 × 領域／工具 ----------
    for tbl_i, fname in [(8, "03_三層×議題領域.csv"), (9, "03_三層×跨領域工具.csv")]:
        src = pd.read_csv(C.TABLE_DIR / fname)
        piv = src.pivot(index="選項", columns="分群", values="分子")
        for row in d.tables[tbl_i].rows[1:]:
            label = row.cells[0].text.strip()
            full = MATCH.get(label, label)
            hit = [x for x in piv.index if x.startswith(full[:8])]
            if not hit:
                continue
            vals = piv.loc[hit[0]]
            for ci, layer in [(1, C.L_NEVER), (2, C.L_AWARE), (3, C.L_DONE)]:
                set_cell(row.cells[ci], f"{vals[layer] / layer_n[layer] * 100:.0f}",
                         f"表{tbl_i} {label} {layer}", f"分母＝{layer_n[layer]}")

    # ---------- 表10：完整狩野附錄 ----------
    all_rows = pd.concat([
        pd.read_csv(C.TABLE_DIR / "02_狩野_全體.csv").assign(分群="全體"),
        kano_lyr,
    ])
    idx = {(r["主題"], r["分群"]): r for _, r in all_rows.iterrows()}
    current_topic = None
    t10_written = []          # (主題, 群體, 寫入的 SI) — 供收尾檢查
    for row in d.tables[10].rows[1:]:
        label = row.cells[0].text.strip()
        stripped = label.lstrip("①②③④⑤")
        if stripped:
            # ⚠ 這裡絕不可用 MATCH.get(stripped, current_topic)：
            #   找不到對應時會靜默沿用上一個主題，把後面四個主題的數字
            #   全部寫成第一個主題的值（曾經真的發生過）。找不到就直接炸掉。
            if stripped not in MATCH:
                raise KeyError(f"表10 主題「{stripped}」在 MATCH 中沒有對應，請先補上再重跑")
            current_topic = MATCH[stripped]
        group = row.cells[1].text.strip()
        key = (current_topic, group)
        if key not in idx:
            continue
        r = idx[key]
        for ci, val in [(2, int(r["A"])), (3, int(r["O"])), (4, int(r["M"])), (5, int(r["I"])),
                        (6, int(r["R"])), (7, int(r["Q"])), (8, int(r["有效樣本"])),
                        (9, f"{r['SI']:.3f}"), (10, f"−{abs(r['DSI']):.3f}"),
                        (11, r["落點"].replace(" ", ""))]:   # 落點象限也要更新，否則會與 SI/DSI 打架
            set_cell(row.cells[ci], val, f"表10 {label}/{group} 第{ci}欄")
        t10_written.append((current_topic, group, f"{r['SI']:.3f}"))

    # 收尾檢查：每一列寫入的 SI，必須等於「該列自己那個主題」在來源表中的值。
    # 不能只檢查「五個值互不相同」——沒資料與捲動的 SI 恰好都是 0.624，
    # 那樣的檢查會誤殺；而真正要防的是「整欄被寫成同一個主題」。
    assert len({t for t, g, _ in t10_written if g == "全體"}) == 5, \
        "表10 的全體列沒有涵蓋五個不同主題——主題對應可能又靜默沿用了"
    for topic, group, si in t10_written:
        expected = idx[(topic, group)]
        assert si == f"{expected['SI']:.3f}", \
            f"表10「{topic}／{group}」寫入 SI={si}，來源表為 {expected['SI']:.3f}"
    print(f"  ✓ 表10 檢查通過：{len(t10_written)} 列各自對應正確主題")

    d.save(DST)

    # ---------- 9. 換掉內嵌的圖 3-1（原本仍是 N=120 版）----------
    # 表格數字更新到 N=126 後，若圖還是舊的，同一份文件會圖表打架。
    # 直接在 zip 裡替換圖片位元組，版面配置與尺寸完全不動（新圖已鎖定同樣的 1560x1118）。
    replace_embedded_figure(
        DST,
        old_png=C.OUT_DIR / "圖表" / "狩野四象限說明圖.png",
        new_png=C.FIG_DIR / "00_狩野四象限說明圖_N126.png",
    )

    ch = pd.DataFrame(CHANGES)
    C.save_table(ch, "08_舊報告變更紀錄", f"共 {len(ch)} 處")
    print(f"\nv2 已輸出：{DST.name}")
    print(f"共修改 {len(ch)} 處，逐筆紀錄在 分析結果/08_舊報告變更紀錄.csv")
    pd.set_option("display.width", 200, "display.max_colwidth", 40)
    print("\n=== 段落層級的變更（文字敘述）===")
    print(ch[ch["位置"].str.startswith("段落")].to_string(index=False))


# 舊報告的簡稱 → 問卷完整選項字串
MATCH = {
    "讓成果落地": "怎麼讓做好的東西真正落地、接進體制",
    "讓成果落地、接進體制": "怎麼讓做好的東西真正落地、接進體制",
    "留住夥伴": "怎麼留住夥伴、維持團隊運作的能量",
    "怎麼留住夥伴、維持團隊能量": "怎麼留住夥伴、維持團隊運作的能量",
    "沒資料怎麼辦": "沒有現成資料時，怎麼搜尋、整理或自建資料集",
    "沒資料時怎麼搜尋、自建資料集": "沒有現成資料時，怎麼搜尋、整理或自建資料集",
    "捲動更多人": "怎麼捲動更多人（包含不會寫程式的人）一起參與",
    # 表10 的列標為「②沒資料」「③捲動」「④留住」「⑤落地」，去掉編號後是這四個短鍵，
    # 缺了它們就會觸發上面的 KeyError（先前正是因為靜默沿用而寫錯 118 處）
    "沒資料": "沒有現成資料時，怎麼搜尋、整理或自建資料集",
    "捲動": "怎麼捲動更多人（包含不會寫程式的人）一起參與",
    "留住": "怎麼留住夥伴、維持團隊運作的能量",
    "落地": "怎麼讓做好的東西真正落地、接進體制",
    "怎麼捲動更多人（含非技術者）": "怎麼捲動更多人（包含不會寫程式的人）一起參與",
    "如何發起": "如何發起一個公民科技專案",
    "如何發起一個公民科技專案": "如何發起一個公民科技專案",
    "政府預算、決算與標案": "政府預算、決算與標案（看政府的錢怎麼花）",
    "群眾協力與通報工具": "群眾協力與通報工具（讓沒有技術背景的人，也能簡單回報、標記，一起累積資料）",
    "資料視覺化工具": "資料視覺化工具",
    "AI 應用（整理、查證資料）": "AI 應用（用 AI 整理、查證資料等）",
    "AI 應用": "AI 應用（用 AI 整理、查證資料等）",
    "地理資訊與地圖工具（GIS）": "地理資訊與地圖工具（GIS、線上地圖）",
    "民意蒐集與線上審議工具": "民意蒐集與線上審議工具（例如問卷、Polis、討論平台）",
    "地方發展與地方創生": "地方發展與地方創生",
    "假訊息與詐騙防治": "假訊息與詐騙防治",
    "法律、立法歷程": "法律、立法歷程",
    "防災與災害應變": "防災與災害應變",
    "怎麼找到現成的研究成果": "怎麼找到現成的研究成果（學術論文、政府委辦研究、考察報告）",
    "開放協作知識庫（Wiki 等）": "開放協作知識庫（維基百科、Wikidata 等共筆知識）",
}
# 完整選項 → 表格用簡稱（回寫表3 時用）。
# 同一個完整選項可能對應多個簡稱，取**最長**的那個，才不會把舊報告的
#「AI 應用（整理、查證資料）」降級成「AI 應用」，無謂地改動 Claire 的用字。
SHORT = {}
for _k, _v in MATCH.items():
    if len(_k) < len(_v) and len(_k) > len(SHORT.get(_v, "")):
        SHORT[_v] = _k

if __name__ == "__main__":
    main()
